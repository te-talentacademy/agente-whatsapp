"""La libreta multi-formato: PDF (texto y escaneado), Word, Excel.

Todo el camino real: archivo en disco -> extracción -> índice -> búsqueda.
El motor de visión se sustituye por dobles (jamás internet); los archivos de
prueba se generan aquí mismo, no hay binarios guardados en el repositorio.
"""

import threading
import time
import zipfile

import pytest

from src import db
from src.llm.client import LlmResult
from src.rag import extract, index

# ---------------------------------------------------------------------------
# Constructores de archivos de prueba
# ---------------------------------------------------------------------------


def _build_pdf(objects: list[bytes]) -> bytes:
    """Arma un PDF mínimo válido con su tabla de posiciones correcta."""
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF".encode()
    return bytes(out)


def make_text_pdf(text: str) -> bytes:
    """Una página carta con el texto dado (fuente Helvetica estándar)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    return _build_pdf([
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ])


def _image_page_objects() -> list[bytes]:
    pixels = bytes([200, 30, 30] * 16)  # imagen roja de 4x4, RGB crudo
    image = (
        b"<< /Type /XObject /Subtype /Image /Width 4 /Height 4 /ColorSpace /DeviceRGB"
        b" /BitsPerComponent 8 /Length " + str(len(pixels)).encode() + b" >>\nstream\n"
        + pixels + b"\nendstream"
    )
    stream = b"q 200 0 0 100 50 600 cm /Im1 Do Q"
    content = b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
    return [image, content]


def make_scanned_pdf(pages: int = 1, mediabox: str = "0 0 612 792") -> bytes:
    """PDF cuyas páginas son SOLO una imagen (sin capa de texto): un escaneo."""
    kids = " ".join(f"{3 + i} 0 R" for i in range(pages)).encode()
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(pages).encode() + b" >>",
    ]
    image_number = 3 + pages
    content_number = image_number + 1
    for i in range(pages):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [{mediabox}] "
            f"/Resources << /XObject << /Im1 {image_number} 0 R >> >> "
            f"/Contents {content_number} 0 R >>".encode()
        )
    objects.extend(_image_page_objects())
    return _build_pdf(objects)


def make_mixed_pdf(text: str) -> bytes:
    """Página 1 con texto + página 2 escaneada (solo imagen)."""
    stream_text = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    pixels = bytes([200, 30, 30] * 16)
    stream_img = b"q 200 0 0 100 50 600 cm /Im1 Do Q"
    return _build_pdf([
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 6 0 R >>"
        ),
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /XObject << /Im1 7 0 R >> >> /Contents 8 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream_text)).encode() + b" >>\nstream\n" + stream_text + b"\nendstream",
        (
            b"<< /Type /XObject /Subtype /Image /Width 4 /Height 4 /ColorSpace /DeviceRGB"
            b" /BitsPerComponent 8 /Length " + str(len(pixels)).encode() + b" >>\nstream\n"
            + pixels + b"\nendstream"
        ),
        b"<< /Length " + str(len(stream_img)).encode() + b" >>\nstream\n" + stream_img + b"\nendstream",
    ])


def make_docx(path, paragraphs: list[str], table: list[list[str]] | None = None) -> None:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    document.save(str(path))


def make_xlsx(path, sheets: dict[str, list[list[object]]]) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    workbook.save(str(path))


# ---------------------------------------------------------------------------
# Apoyos
# ---------------------------------------------------------------------------


@pytest.fixture()
def knowledge_dir(tmp_path, monkeypatch):
    folder = tmp_path / "conocimiento"
    folder.mkdir()
    monkeypatch.setattr("src.config.KNOWLEDGE_DIR", str(folder))
    monkeypatch.setenv("FEATURE_RAG", "on")
    return folder


class FakeVision:
    """Doble del motor de visión: respuestas programadas + conteo de llamadas."""

    def __init__(self, monkeypatch, result: LlmResult | None = None, delay: float = 0.0):
        self.calls = 0
        self.result = result or LlmResult(ok=True, text="La tarta de fresa cuesta 90 pesos")
        self.delay = delay
        monkeypatch.setenv("FEATURE_VISION", "on")
        monkeypatch.setenv("OPENROUTER_API_KEY", "llave-de-pruebas")
        monkeypatch.setattr("src.llm.client.complete", self._complete)

    def _complete(self, messages, model, api_key, title="", max_tokens=None, temperature=None):
        self.calls += 1
        self.last_max_tokens = max_tokens
        self.last_temperature = temperature
        if self.delay:
            time.sleep(self.delay)
        return self.result


def _rebuild():
    return index.rebuild()


def _hits(question: str) -> list[dict]:
    return index.search(question)


# ---------------------------------------------------------------------------
# Formatos: el camino completo archivo -> índice -> búsqueda
# ---------------------------------------------------------------------------


def test_pdf_con_texto_se_indexa_y_responde(knowledge_dir):
    (knowledge_dir / "precios.pdf").write_bytes(
        make_text_pdf("Los croissants cuestan 30 pesos cada uno")
    )
    assert _rebuild() > 0
    hits = _hits("cuanto cuestan los croissants")
    assert hits and "croissants" in hits[0]["chunk"]


def test_docx_con_parrafos_y_tabla(knowledge_dir):
    make_docx(
        knowledge_dir / "politicas.docx",
        ["Los pedidos grandes requieren dos dias de anticipo."],
        table=[["Producto", "Precio"], ["Rosca de almendra", "180 pesos"]],
    )
    assert _rebuild() > 0
    assert any("anticipo" in h["chunk"] for h in _hits("pedidos grandes anticipo"))
    assert any("180" in h["chunk"] for h in _hits("precio rosca almendra"))


def test_xlsx_con_hojas_y_encabezado(knowledge_dir):
    make_xlsx(
        knowledge_dir / "inventario.xlsx",
        {"Precios": [["Concha de vainilla", 12], ["Bolillo", 4]]},
    )
    assert _rebuild() > 0
    hits = _hits("precio concha vainilla")
    assert hits and "Concha de vainilla" in hits[0]["chunk"]
    assert any("Hoja: Precios" in h["chunk"] for h in hits)


def test_md_y_txt_siguen_igual(knowledge_dir):
    (knowledge_dir / "notas.md").write_text("# Horarios\n\nAbrimos a las 8 de la manana.")
    (knowledge_dir / "extra.txt").write_text("Los sabados cerramos a las 2.")
    assert _rebuild() > 0
    assert any("8" in h["chunk"] for h in _hits("a que hora abrimos"))
    assert any("cerramos" in h["chunk"] for h in _hits("los sabados cerramos"))


def test_pdf_escaneado_pasa_por_vision_y_cachea(knowledge_dir, monkeypatch):
    vision = FakeVision(monkeypatch)
    (knowledge_dir / "menu.pdf").write_bytes(make_scanned_pdf())
    assert _rebuild() > 0
    assert vision.calls == 1
    assert vision.last_temperature == 0.0
    assert vision.last_max_tokens == extract.OCR_MAX_TOKENS
    assert any("tarta de fresa" in h["chunk"] for h in _hits("cuanto cuesta la tarta de fresa"))
    # Segundo rearme: misma página servida de la memoria, cero llamadas nuevas.
    assert _rebuild() > 0
    assert vision.calls == 1


# ---------------------------------------------------------------------------
# Fallos y preservación del índice
# ---------------------------------------------------------------------------


def test_archivo_corrupto_se_salta_y_el_resto_se_indexa(knowledge_dir):
    (knowledge_dir / "roto.pdf").write_bytes(b"esto no es un pdf")
    (knowledge_dir / "roto.docx").write_bytes(b"tampoco es un zip")
    (knowledge_dir / "sano.txt").write_text("El pan de muerto sale en octubre.")
    assert _rebuild() > 0
    assert any("octubre" in h["chunk"] for h in _hits("pan de muerto"))


def test_version_corrupta_conserva_el_conocimiento_anterior(knowledge_dir):
    path = knowledge_dir / "manual.docx"
    make_docx(path, ["El telefono de pedidos es 5550101."])
    assert _rebuild() > 0
    assert any("5550101" in h["chunk"] for h in _hits("telefono de pedidos"))
    # Se sube por accidente una versión rota CON EL MISMO NOMBRE.
    path.write_bytes(b"bytes de un archivo danado")
    assert _rebuild() > 0
    assert any("5550101" in h["chunk"] for h in _hits("telefono de pedidos"))


def test_fallo_pasajero_de_vision_conserva_la_version_anterior(knowledge_dir, monkeypatch):
    vision = FakeVision(monkeypatch)
    path = knowledge_dir / "menu.pdf"
    path.write_bytes(make_scanned_pdf())
    assert _rebuild() > 0
    # El archivo cambia (hash nuevo) y ahora la visión falla de forma pasajera.
    path.write_bytes(make_scanned_pdf(mediabox="0 0 611 791"))
    vision.result = LlmResult(ok=False, retryable=True, reason="servicio saturado (429)")
    monkeypatch.setattr(extract, "OCR_RETRY_PAUSE_SECONDS", 0.0)
    assert _rebuild() > 0
    assert any("tarta de fresa" in h["chunk"] for h in _hits("tarta de fresa"))


def test_respuesta_truncada_es_pagina_fallida_no_cacheada(knowledge_dir, monkeypatch):
    vision = FakeVision(monkeypatch, result=LlmResult(ok=True, text="parcial...", truncated=True))
    (knowledge_dir / "menu.pdf").write_bytes(make_scanned_pdf())
    _rebuild()
    assert vision.calls >= 1
    assert not _hits("parcial")
    with db.transaction() as conn:
        cached = conn.execute("SELECT COUNT(*) AS c FROM ocr_cache").fetchone()["c"]
    assert cached == 0


def test_archivo_eliminado_sale_del_indice(knowledge_dir):
    path = knowledge_dir / "viejo.txt"
    path.write_text("La promocion de enero ya termino.")
    assert _rebuild() > 0
    assert _hits("promocion de enero")
    path.unlink()
    _rebuild()
    assert not _hits("promocion de enero")


# ---------------------------------------------------------------------------
# Visión apagada
# ---------------------------------------------------------------------------


def test_escaneado_sin_vision_se_salta_con_claridad(knowledge_dir, monkeypatch):
    monkeypatch.delenv("FEATURE_VISION", raising=False)
    (knowledge_dir / "menu.pdf").write_bytes(make_scanned_pdf())
    (knowledge_dir / "sano.txt").write_text("Hacemos entregas a domicilio.")
    stats = extract.RebuildStats()
    result = extract.extract(str(knowledge_dir / "menu.pdf"), stats)
    assert not result.ok and "visión" in result.reason
    assert _rebuild() > 0
    assert any("domicilio" in h["chunk"] for h in _hits("entregas a domicilio"))


def test_paginas_ya_cacheadas_se_reusan_con_vision_apagada(knowledge_dir, monkeypatch):
    vision = FakeVision(monkeypatch)
    (knowledge_dir / "menu.pdf").write_bytes(make_scanned_pdf())
    assert _rebuild() > 0
    assert vision.calls == 1
    # La visión se apaga: la página ya pagada sigue sirviendo desde la memoria.
    monkeypatch.delenv("FEATURE_VISION", raising=False)
    assert _rebuild() > 0
    assert vision.calls == 1
    assert any("tarta de fresa" in h["chunk"] for h in _hits("tarta de fresa"))


# ---------------------------------------------------------------------------
# Topes y presupuestos
# ---------------------------------------------------------------------------


def test_tope_por_documento_publica_lo_leido(knowledge_dir, monkeypatch):
    vision = FakeVision(monkeypatch)
    monkeypatch.setattr(extract, "SCANNED_PAGES_PER_DOC", 1)
    (knowledge_dir / "menu.pdf").write_bytes(make_scanned_pdf(pages=2))
    stats = extract.RebuildStats()
    result = extract.extract(str(knowledge_dir / "menu.pdf"), stats, None)
    assert result.ok and result.skipped_pages == 1
    assert vision.calls == 1


def test_tope_por_documento_estable_entre_rearmes(knowledge_dir, monkeypatch):
    """El acierto de memoria también avanza el ordinal: la página 2 jamás se
    cuela en un despliegue posterior porque la página 1 ya esté pagada."""
    vision = FakeVision(monkeypatch)
    monkeypatch.setattr(extract, "SCANNED_PAGES_PER_DOC", 1)
    (knowledge_dir / "menu.pdf").write_bytes(make_scanned_pdf(pages=2))
    first = extract.RebuildStats()
    r1 = extract.extract(str(knowledge_dir / "menu.pdf"), first, None)
    assert r1.ok and r1.skipped_pages == 1 and vision.calls == 1
    second = extract.RebuildStats()
    r2 = extract.extract(str(knowledge_dir / "menu.pdf"), second, None)
    assert r2.ok and r2.skipped_pages == 1
    assert vision.calls == 1          # histórico total: sigue siendo UNA llamada
    assert second.cache_hits == 1     # la página 1 vino de la memoria
    assert second.http_calls == 0


def test_techo_global_cuenta_reintentos(knowledge_dir, monkeypatch):
    vision = FakeVision(
        monkeypatch, result=LlmResult(ok=False, retryable=True, reason="saturado")
    )
    monkeypatch.setattr(extract, "OCR_RETRY_PAUSE_SECONDS", 0.0)
    (knowledge_dir / "menu.pdf").write_bytes(make_scanned_pdf())
    stats = extract.RebuildStats()
    extract.extract(str(knowledge_dir / "menu.pdf"), stats, None)
    # Una página con 2 intentos = 2 solicitudes reales contra el techo.
    assert vision.calls == 2
    assert stats.http_calls == 2


def test_techo_global_agotado_deja_pendiente_sin_publicar_a_medias(knowledge_dir, monkeypatch):
    vision = FakeVision(monkeypatch)
    monkeypatch.setattr(extract, "SCANNED_PAGES_PER_REBUILD", 0)
    (knowledge_dir / "menu.pdf").write_bytes(make_scanned_pdf())
    stats = extract.RebuildStats()
    result = extract.extract(str(knowledge_dir / "menu.pdf"), stats, None)
    assert not result.ok and "techo" in result.reason
    assert stats.pending_pages == 1
    assert vision.calls == 0


def test_mediabox_absurdo_se_salta_sin_dibujar(knowledge_dir, monkeypatch):
    vision = FakeVision(monkeypatch)
    (knowledge_dir / "raro.pdf").write_bytes(make_scanned_pdf(mediabox="0 0 200000 200000"))
    stats = extract.RebuildStats()
    result = extract.extract(str(knowledge_dir / "raro.pdf"), stats, None)
    # La única página tiene medidas fuera de rango: se salta ANTES de crear la
    # imagen, sin una sola llamada de visión, y el documento (que queda sin
    # texto) termina en ok=False con motivo claro.
    assert result.ok is False
    assert "legible" in result.reason
    assert vision.calls == 0
    assert stats.http_calls == 0


def test_presupuesto_de_pixeles_reduce_la_escala(knowledge_dir, monkeypatch):
    vision = FakeVision(monkeypatch)
    # Página A2 aproximada (1191 x 1684 puntos): a 150 dpi excede los 4 MP,
    # pero cabe reduciendo la escala — se lee igual, sin reventar la memoria.
    (knowledge_dir / "poster.pdf").write_bytes(make_scanned_pdf(mediabox="0 0 1191 1684"))
    stats = extract.RebuildStats()
    result = extract.extract(str(knowledge_dir / "poster.pdf"), stats, None)
    assert result.ok and result.used_vision
    assert vision.calls == 1


def test_zip_bomba_se_rechaza_antes_de_abrir(knowledge_dir):
    path = knowledge_dir / "bomba.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\x00" * (60 * 1024 * 1024))
    stats = extract.RebuildStats()
    result = extract.extract(str(path), stats)
    assert not result.ok and "descomprimir" in result.reason


def test_hoja_gigante_se_rechaza_por_celdas(knowledge_dir, monkeypatch):
    monkeypatch.setattr(extract, "MAX_SHEET_CELLS", 10)
    make_xlsx(knowledge_dir / "grande.xlsx", {"Datos": [[1, 2, 3, 4] for _ in range(5)]})
    stats = extract.RebuildStats()
    result = extract.extract(str(knowledge_dir / "grande.xlsx"), stats)
    assert not result.ok and "celdas" in result.reason


def test_archivo_demasiado_pesado_se_salta(knowledge_dir, monkeypatch):
    monkeypatch.setattr(extract, "MAX_FILE_BYTES", 100)
    (knowledge_dir / "pesado.pdf").write_bytes(b"%" * 200)
    stats = extract.RebuildStats()
    result = extract.extract(str(knowledge_dir / "pesado.pdf"), stats)
    assert not result.ok and "MB" in result.reason


def test_presupuesto_de_texto_corta_durante_la_extraccion(knowledge_dir, monkeypatch):
    """El corte es INCREMENTAL: al rebasar, ni se sigue construyendo el texto
    ni se visitan las páginas posteriores (la escaneada jamás llega a visión)."""
    vision = FakeVision(monkeypatch)
    monkeypatch.setattr(extract, "MAX_TEXT_CHARS", 10)
    (knowledge_dir / "mixto.pdf").write_bytes(
        make_mixed_pdf("Los croissants cuestan 30 pesos cada uno")
    )
    stats = extract.RebuildStats()
    result = extract.extract(str(knowledge_dir / "mixto.pdf"), stats, None)
    assert not result.ok and "grande" in result.reason
    assert vision.calls == 0
    assert stats.http_calls == 0


def test_presupuesto_de_texto_en_excel(knowledge_dir, monkeypatch):
    monkeypatch.setattr(extract, "MAX_TEXT_CHARS", 20)
    make_xlsx(
        knowledge_dir / "grande.xlsx",
        {"Datos": [["fila con texto bastante largo"] for _ in range(10)]},
    )
    stats = extract.RebuildStats()
    result = extract.extract(str(knowledge_dir / "grande.xlsx"), stats)
    assert not result.ok and "grande" in result.reason


# ---------------------------------------------------------------------------
# Concurrencia y apagado (contratos sellados en la revisión del plan)
# ---------------------------------------------------------------------------


def test_ocr_lento_no_bloquea_al_timbre(knowledge_dir, monkeypatch):
    FakeVision(monkeypatch, delay=1.2)
    (knowledge_dir / "menu.pdf").write_bytes(make_scanned_pdf())
    worker = threading.Thread(target=_rebuild, daemon=True)
    worker.start()
    time.sleep(0.3)  # el rearme ya está dentro de la llamada lenta de visión
    from src.queue import store
    from src.webhook.parse import InboundMessage

    started = time.monotonic()
    store.enqueue(InboundMessage(wamid="wamid.CONCURRENTE", sender="521", kind="text", body="hola", timestamp="1"))
    elapsed = time.monotonic() - started
    worker.join(timeout=10)
    # La escritura del timbre no espera a la visión: mucho menos que el candado.
    assert elapsed < 1.0


def test_apagado_cooperativo_corta_sin_empezar_otra_pagina(knowledge_dir, monkeypatch):
    class SignalDuringCall(FakeVision):
        def _complete(self, *args, **kwargs):
            result = super()._complete(*args, **kwargs)
            index.STOP_EVENT.set()   # la señal llega mientras se lee la página 1
            return result

    vision = SignalDuringCall(monkeypatch)
    (knowledge_dir / "menu.pdf").write_bytes(make_scanned_pdf(pages=3))
    total = _rebuild()
    # Tras la señal no empieza otra página, ni otra solicitud, ni NINGUNA
    # transacción de memoria (la página en curso tampoco se guarda), ni la
    # publicación.
    assert vision.calls == 1
    assert total == 0
    with db.transaction() as conn:
        assert conn.execute("SELECT COUNT(*) AS c FROM ocr_cache").fetchone()["c"] == 0
    assert not _hits("tarta de fresa")


def test_apagado_durante_reintento_ni_gasta_ni_cachea(knowledge_dir, monkeypatch):
    """La señal se consulta DENTRO del ciclo de reintentos y la pausa la
    observa: exactamente una solicitud, cero memoria, cero publicación."""

    class SignalOnFailure(FakeVision):
        def _complete(self, *args, **kwargs):
            result = super()._complete(*args, **kwargs)
            index.STOP_EVENT.set()   # la señal llega mientras el intento 1 falla
            return result

    vision = SignalOnFailure(
        monkeypatch, result=LlmResult(ok=False, retryable=True, reason="saturado")
    )
    monkeypatch.setattr(extract, "OCR_RETRY_PAUSE_SECONDS", 30.0)
    (knowledge_dir / "menu.pdf").write_bytes(make_scanned_pdf())
    started = time.monotonic()
    total = _rebuild()
    elapsed = time.monotonic() - started
    assert vision.calls == 1   # el segundo intento jamás sale
    assert total == 0          # nada se publica
    with db.transaction() as conn:
        assert conn.execute("SELECT COUNT(*) AS c FROM ocr_cache").fetchone()["c"] == 0
    assert elapsed < 5.0       # la pausa de 30 s observó la señal y cortó al instante


def test_lifespan_captura_el_fallo_del_rearme_y_conserva_el_indice(knowledge_dir, monkeypatch, caplog):
    (knowledge_dir / "sano.txt").write_text("Aceptamos pagos con tarjeta.")
    assert _rebuild() > 0

    def boom():
        raise RuntimeError("fallo inesperado del extractor")

    monkeypatch.setattr(index, "rebuild", boom)
    import logging

    from fastapi.testclient import TestClient

    from src.main import app

    with caplog.at_level(logging.ERROR, logger="agente"):
        with TestClient(app) as web:
            deadline = time.monotonic() + 5.0
            while "No pude armar la libreta" not in caplog.text and time.monotonic() < deadline:
                time.sleep(0.05)
            assert web.get("/").status_code == 200  # el servicio sigue de pie
    assert "No pude armar la libreta" in caplog.text
    assert any("tarjeta" in h["chunk"] for h in _hits("pagos con tarjeta"))
